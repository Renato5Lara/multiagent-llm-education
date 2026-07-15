"""
Exportación de resultados experimentales para análisis estadístico
(SPSS / RStudio / Python).

CSV: stdlib + BOM UTF-8 (utf-8-sig) para compatibilidad SPSS/Excel.
XLSX: openpyxl (import guard — si falta, el CSV sigue funcionando).
"""

import csv
import io
import logging

logger = logging.getLogger(__name__)

# Orden y nombres de columnas SPSS-safe (sin espacios ni tildes)
EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("student_id", "id_estudiante"),
    ("group", "grupo"),
    ("pre_pct", "pretest_pct"),
    ("post_pct", "posttest_pct"),
    ("absolute_gain", "incremento_absoluto"),
    ("normalized_gain", "ganancia_normalizada"),
    ("level", "nivel"),
    ("pre_level", "nivel_pre"),
    ("post_level", "nivel_post"),
    ("path_generation_ms", "tiempo_ruta_ms"),
    ("ai_time_ms", "tiempo_ia_ms"),
    ("total_time_seconds", "tiempo_total_seg"),
    ("profile", "perfil"),
    ("course", "curso"),
    ("date", "fecha"),
]

# Detalle por ciclo (submit_cycle_evidence) — modalidad diagnosticada vs.
# modalidad de refuerzo, la comparación central de la hipótesis de
# adaptación multimodal.
CYCLE_EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("student_id", "id_estudiante"),
    ("email", "correo"),
    ("concepto", "concepto"),
    ("modalidad_diagnosticada", "modalidad_diagnosticada"),
    ("modalidad_refuerzo", "modalidad_refuerzo"),
    ("profundidad", "profundidad"),
    ("intentos", "intentos"),
    ("resultado", "resultado"),
    ("ayudas", "ayudas_utilizadas"),
    ("tiempo_ms", "tiempo_ms"),
    ("fecha", "fecha"),
]


def excel_available() -> bool:
    try:
        import openpyxl  # noqa: F401

        return True
    except ImportError:
        return False


def rows_to_csv(rows: list[dict]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([header for _, header in EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([_cell(row.get(key)) for key, _ in EXPORT_COLUMNS])
    # BOM UTF-8 para que SPSS/Excel detecten la codificación
    return "\ufeff" + buffer.getvalue()


def rows_to_xlsx(rows: list[dict], summary: dict | None = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append([header for _, header in EXPORT_COLUMNS])
    for row in rows:
        ws.append([_cell(row.get(key)) for key, _ in EXPORT_COLUMNS])

    if summary:
        ws2 = wb.create_sheet("Resumen")
        ws2.append(["metrica", "valor"])
        for key, value in summary.items():
            if isinstance(value, dict):
                for sub_key, sub_value in value.items():
                    ws2.append([f"{key}.{sub_key}", _cell(sub_value)])
            else:
                ws2.append([key, _cell(value)])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _write_summary_sheet(ws, summary: dict) -> None:
    ws.append(["metrica", "valor"])
    for key, value in summary.items():
        if isinstance(value, dict):
            for sub_key, sub_value in value.items():
                ws.append([f"{key}.{sub_key}", _cell(sub_value)])
        else:
            ws.append([key, _cell(value)])


def experiment_to_xlsx(
    summary_rows: list[dict], cycle_rows: list[dict], statistics: dict
) -> bytes:
    """Exportar experimento — un solo archivo con las tres vistas que la
    investigación necesita: resumen por estudiante (pre/post/ganancia/
    perfil), detalle por ciclo (modalidad diagnosticada vs. modalidad de
    refuerzo real) y estadísticas descriptivas. Reutiliza exactamente los
    mismos datos que ya sirven /students y /export — ningún cálculo nuevo,
    solo compuestos en un único archivo con un clic."""
    from openpyxl import Workbook

    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen.append([header for _, header in EXPORT_COLUMNS])
    for row in summary_rows:
        ws_resumen.append([_cell(row.get(key)) for key, _ in EXPORT_COLUMNS])

    ws_ciclos = wb.create_sheet("Ciclos")
    ws_ciclos.append([header for _, header in CYCLE_EXPORT_COLUMNS])
    for row in cycle_rows:
        ws_ciclos.append([_cell(row.get(key)) for key, _ in CYCLE_EXPORT_COLUMNS])

    ws_stats = wb.create_sheet("Estadisticas")
    _write_summary_sheet(ws_stats, statistics)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _cell(value):
    if value is None:
        return ""
    return value
