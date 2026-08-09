"""
Tests del Dashboard del Investigador y la exportación CSV/Excel.
"""

import csv
import io

from tests.conftest import auth_header

from app.data.knowledge_test_bank import seed_knowledge_test_bank
from app.models.knowledge_test import KnowledgeTestQuestion
from app.services import research_export_service


def _complete_pre_and_post(client, token, course_id, db, student_id, pre_correct_ratio=0.0):
    """Rinde pre (con la fracción de aciertos dada) y post (todo correcto).

    El Post-Test exige una Ruta de Aprendizaje completa (recorrido real): se
    simula que el estudiante ya terminó sus módulos antes de rendirlo.
    """
    bank = {q.id: q.correct_index for q in db.query(KnowledgeTestQuestion).all()}

    start = client.post(
        f"/api/students/knowledge-test/{course_id}/start",
        headers=auth_header(token),
        json={"kind": "pre"},
    ).json()
    n_correct = int(len(start["questions"]) * pre_correct_ratio)
    answers = {}
    for i, q in enumerate(start["questions"]):
        correct = bank[q["id"]]
        answers[q["id"]] = correct if i < n_correct else (correct + 1) % 4
    client.post(
        f"/api/students/knowledge-test/attempt/{start['attempt_id']}/submit",
        headers=auth_header(token),
        json={"answers": answers},
    )

    from app.models.student_progress import LearningPath

    db.add(
        LearningPath(
            student_id=student_id,
            course_id=course_id,
            total_modules=2,
            status="active",
        )
    )
    db.commit()

    start = client.post(
        f"/api/students/knowledge-test/{course_id}/start",
        headers=auth_header(token),
        json={"kind": "post"},
    ).json()
    answers = {q["id"]: bank[q["id"]] for q in start["questions"]}
    client.post(
        f"/api/students/knowledge-test/attempt/{start['attempt_id']}/submit",
        headers=auth_header(token),
        json={"answers": answers},
    )


def test_summary_empty_database(client, db, admin_token):
    resp = client.get("/api/research/summary", headers=auth_header(admin_token))
    assert resp.status_code == 200
    body = resp.json()
    assert body["n_students_pretested"] == 0
    assert body["avg_pre_pct"] is None
    assert body["group_label"] == "Experimental"


def test_summary_requires_admin_or_docente_role(client, estudiante_token):
    resp = client.get("/api/research/summary", headers=auth_header(estudiante_token))
    assert resp.status_code == 403


def test_summary_rejects_unauthenticated(client):
    resp = client.get("/api/research/summary")
    assert resp.status_code == 401


def test_summary_reflects_real_attempts(
    client, estudiante_token, admin_token, curso_publicado, db, estudiante_user
):
    seed_knowledge_test_bank(db)
    _complete_pre_and_post(client, estudiante_token, curso_publicado.id, db, estudiante_user.id)

    body = client.get("/api/research/summary", headers=auth_header(admin_token)).json()
    assert body["n_students_pretested"] == 1
    assert body["n_students_posttested"] == 1
    assert body["n_compared"] == 1
    assert body["avg_pre_pct"] == 0.0
    assert body["avg_post_pct"] == 100.0
    assert body["avg_absolute_gain"] == 100.0
    assert body["level_distribution_pre"]["basico"] == 1
    assert body["level_distribution_post"]["avanzado"] == 1


def test_students_rows_dataset(
    client, estudiante_token, admin_token, curso_publicado, db, estudiante_user
):
    seed_knowledge_test_bank(db)
    _complete_pre_and_post(client, estudiante_token, curso_publicado.id, db, estudiante_user.id)

    body = client.get("/api/research/students", headers=auth_header(admin_token)).json()
    assert body["total"] == 1
    row = body["rows"][0]
    assert row["group"] == "Experimental"
    assert row["pre_pct"] == 0.0
    assert row["post_pct"] == 100.0
    assert row["absolute_gain"] == 100.0
    assert row["level"] == "avanzado"
    assert row["course"] == curso_publicado.code
    assert row["date"] is not None


def test_export_csv_is_spss_ready(
    client, estudiante_token, admin_token, curso_publicado, db, estudiante_user
):
    seed_knowledge_test_bank(db)
    _complete_pre_and_post(client, estudiante_token, curso_publicado.id, db, estudiante_user.id)

    resp = client.get("/api/research/export?fmt=csv", headers=auth_header(admin_token))
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    assert "attachment" in resp.headers["content-disposition"]

    text = resp.content.decode("utf-8-sig")  # el BOM debe estar presente
    assert resp.content.startswith("﻿".encode("utf-8"))
    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    assert len(rows) == 1
    assert rows[0]["grupo"] == "Experimental"
    assert rows[0]["pretest_pct"] == "0.0"
    assert rows[0]["posttest_pct"] == "100.0"
    assert set(reader.fieldnames) == {
        header for _, header in research_export_service.EXPORT_COLUMNS
    }


def test_export_xlsx_when_available(
    client, estudiante_token, admin_token, curso_publicado, db, estudiante_user
):
    seed_knowledge_test_bank(db)
    _complete_pre_and_post(client, estudiante_token, curso_publicado.id, db, estudiante_user.id)

    resp = client.get("/api/research/export?fmt=xlsx", headers=auth_header(admin_token))
    assert resp.status_code == 200
    if research_export_service.excel_available():
        assert "spreadsheetml" in resp.headers["content-type"]
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(resp.content))
        assert "Resultados" in wb.sheetnames
        assert "Resumen" in wb.sheetnames
        ws = wb["Resultados"]
        assert ws.max_row == 2  # cabecera + 1 estudiante
    else:
        assert "text/csv" in resp.headers["content-type"]


def test_export_rejects_unknown_format(client, admin_token):
    resp = client.get("/api/research/export?fmt=pdf", headers=auth_header(admin_token))
    assert resp.status_code == 422


def test_export_experiment_has_three_sheets(
    client, estudiante_token, admin_token, curso_publicado, db, estudiante_user
):
    """Exportar experimento: un archivo, tres hojas — resumen, ciclos y
    estadísticas — sin recalcular nada que /students o /export ya cubran."""
    seed_knowledge_test_bank(db)

    # Orden real garantizado por la UI (Dashboard.tsx: !has_diagnostic
    # bloquea "Ver ruta adaptativa"): el diagnóstico Likert siempre existe
    # ANTES del pre-test — de lo contrario, el pre-test es el primer hecho
    # de la sesión y fija la única decisión "reforzar" por sesión con el
    # valor por defecto, no con la modalidad real.
    from app.models.diagnostic_result import DiagnosticResult

    db.add(
        DiagnosticResult(
            student_id=estudiante_user.id,
            course_id=curso_publicado.id,
            answers={},
            profile={},
            modality_scores={"kinesthetic": 5.0},
            dominant_modality="kinesthetic",
        )
    )
    db.commit()

    _complete_pre_and_post(client, estudiante_token, curso_publicado.id, db, estudiante_user.id)

    # attempts>=2 y solved=False fuerza "reforzar" (scoring-v1: errores>=2 ⇒
    # no dominada) — el caso donde Adaptar debe honrar la modalidad
    # diagnosticada; "avanzar-con-andamiaje" usa "mixta" a propósito, que
    # nunca coincide con ninguna de las 4 modalidades VARK.
    resp = client.post(
        "/api/students/cycle-evidence",
        headers=auth_header(estudiante_token),
        json={
            "course_id": curso_publicado.id,
            "competencia": "variables",
            "attempts": 3,
            "solved": False,
            "hints_used": 1,
            "time_ms": 12345,
        },
    )
    assert resp.status_code == 200

    resp = client.get("/api/research/export-experiment", headers=auth_header(admin_token))
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Resumen", "Ciclos", "Estadisticas"]

    ws_resumen = wb["Resumen"]
    assert ws_resumen.max_row == 2  # cabecera + 1 estudiante

    ws_ciclos = wb["Ciclos"]
    assert ws_ciclos.max_row == 2  # cabecera + 1 ciclo
    ciclo_row = dict(zip(
        [c.value for c in ws_ciclos[1]],
        [c.value for c in ws_ciclos[2]],
    ))
    assert ciclo_row["concepto"] == "variables"
    assert ciclo_row["modalidad_diagnosticada"] == "kinesthetic"
    assert ciclo_row["ayudas_utilizadas"] == 1

    ws_stats = wb["Estadisticas"]
    stats = dict(
        (row[0].value, row[1].value) for row in ws_stats.iter_rows(min_row=2)
    )
    assert stats["n_ciclos_registrados"] == 1
    assert stats["porcentaje_coincidencia_modalidad"] in (100, 100.0)
